"""0.9.119 — SDS/PKD kimyasal ürün sicili; 0.9.120 — GHS tehlike etiketi checklist."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
import json

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.api.company_access import company_ids_for_query, ensure_company_access
from app.api.deps import get_current_user, require_roles_or_workplace_operations
from app.core.database import get_db
from app.models.entities import (
    Branch,
    ChemicalProduct,
    Company,
    DocumentCategory,
    DocumentRecord,
    ExplosionProtectionDocument,
    User,
    UserRole,
)
from app.schemas.sds import (
    ChemicalProductCreate,
    ChemicalProductResponse,
    ChemicalProductUpdate,
    GhsChecklistUpdate,
    PkdDocumentCreate,
    PkdDocumentResponse,
    PkdDocumentUpdate,
    PkdSummary,
    SdsDueSummary,
)
from app.services.ghs_label_checklist import (
    GHS_ENGINE,
    catalog as ghs_catalog,
    parse_checklist,
    serialize_selected,
)

router = APIRouter(prefix="/sds", tags=["SDS / PKD"])

EDIT_ROLES = (
    UserRole.GLOBAL_ADMIN,
    UserRole.SAFETY_SPECIALIST,
)
VIEW_ROLES = (
    UserRole.GLOBAL_ADMIN,
    UserRole.SAFETY_SPECIALIST,
    UserRole.WORKPLACE_PHYSICIAN,
    UserRole.OTHER_HEALTH_PERSONNEL,
)

REGISTER_ENGINE = "chemical-register-v1"
PKD_ENGINE = "pkd-register-v1"
PKD_STATUS_LABELS = {
    "draft": "Taslak",
    "active": "Aktif",
    "revision_pending": "Revizyon Bekliyor",
    "archived": "Arşiv",
}
PKD_ATMOSPHERE_LABELS = {
    "gas_vapour_mist": "Gaz / buhar / sis",
    "dust": "Yanıcı toz",
    "mixed": "Gaz ve toz birlikte",
}
PKD_ZONE_LABELS = {
    "zone_0": "Zone 0",
    "zone_1": "Zone 1",
    "zone_2": "Zone 2",
    "zone_20": "Zone 20",
    "zone_21": "Zone 21",
    "zone_22": "Zone 22",
}


def _review_status(next_review: date | None) -> str:
    if not next_review:
        return "unset"
    today = date.today()
    if next_review < today:
        return "overdue"
    if next_review <= today + timedelta(days=30):
        return "due_soon"
    return "ok"


def _to_response(row: ChemicalProduct) -> ChemicalProductResponse:
    ghs = parse_checklist(getattr(row, "ghs_checklist_json", None))
    return ChemicalProductResponse(
        id=row.id,
        company_id=row.company_id,
        branch_id=row.branch_id,
        product_name=row.product_name,
        cas_number=row.cas_number,
        has_sds_file=bool(row.has_sds_file),
        document_id=row.document_id,
        next_review_date=row.next_review_date,
        notes=row.notes,
        is_active=bool(row.is_active),
        created_by_id=row.created_by_id,
        created_at=row.created_at,
        updated_at=row.updated_at,
        review_status=_review_status(row.next_review_date),
        ghs_selected=list(ghs["selected"]),
        ghs_count=int(ghs["count"]),
    )


def _ensure_edit(db: Session, user: User, company_id: int) -> None:
    ensure_company_access(db, user, company_id)


def _validated_branch(db: Session, company_id: int, branch_id: int | None) -> Branch | None:
    if branch_id is None:
        return None
    branch = db.get(Branch, branch_id)
    if not branch or not branch.is_active or int(branch.company_id) != int(company_id):
        raise HTTPException(400, "Seçilen şube bu firmaya ait değil veya aktif değil.")
    return branch


def _json_list(raw: str | None) -> list[str]:
    if not raw:
        return []
    try:
        value = json.loads(raw)
    except (TypeError, ValueError):
        return []
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _serialize_list(values: list[str] | None) -> str:
    return json.dumps(values or [], ensure_ascii=False, separators=(",", ":"))


def _pkd_review_status(row: ExplosionProtectionDocument) -> str:
    if not row.is_active or row.status == "archived":
        return "archived"
    if row.status == "draft":
        return "draft"
    if row.status == "revision_pending":
        return "revision_pending"
    if not row.next_review_date:
        return "unset"
    today = date.today()
    if row.next_review_date < today:
        return "overdue"
    if row.next_review_date <= today + timedelta(days=30):
        return "due_soon"
    return "ok"


def _pkd_to_response(row: ExplosionProtectionDocument) -> PkdDocumentResponse:
    return PkdDocumentResponse(
        id=row.id,
        company_id=row.company_id,
        branch_id=row.branch_id,
        document_no=row.document_no,
        area_name=row.area_name,
        process_name=row.process_name,
        atmosphere_type=row.atmosphere_type,
        hazardous_materials=row.hazardous_materials,
        zone_classifications=_json_list(row.zone_classifications_json),
        ignition_sources=_json_list(row.ignition_sources_json),
        control_measures=_json_list(row.control_measures_json),
        responsible_person=row.responsible_person,
        prepared_by=row.prepared_by,
        approved_by=row.approved_by,
        document_date=row.document_date,
        revision_no=row.revision_no,
        next_review_date=row.next_review_date,
        status=row.status,
        has_pkd_file=bool(row.has_pkd_file),
        document_id=row.document_id,
        notes=row.notes,
        is_active=bool(row.is_active),
        created_by_id=row.created_by_id,
        created_at=row.created_at,
        updated_at=row.updated_at,
        review_status=_pkd_review_status(row),
    )


@router.get("/meta")
def sds_meta(user: User = Depends(get_current_user)):
    return {
        "engine": REGISTER_ENGINE,
        "pkd_engine": PKD_ENGINE,
        "ghs_engine": GHS_ENGINE,
        "version_feature": "0.9.120",
        "fields": [
            "product_name",
            "cas_number",
            "has_sds_file",
            "next_review_date",
            "document_id",
            "ghs_checklist",
        ],
        "ghs_pictograms": ghs_catalog(),
        "pkd_statuses": PKD_STATUS_LABELS,
        "pkd_atmosphere_types": PKD_ATMOSPHERE_LABELS,
        "pkd_zones": PKD_ZONE_LABELS,
        "note": "SDS ve PKD dosyaları Dokümanlar kaydı üzerinden yüklenir; PKD kayıtları işyeri/bölüm/proses bazında izlenir.",
    }


@router.get("/ghs-catalog")
def ghs_label_catalog(user: User = Depends(get_current_user)):
    return {"engine": GHS_ENGINE, "pictograms": ghs_catalog()}


@router.get("/export.xlsx")
def export_sds_xlsx(
    company_id: int | None = None,
    q: str | None = Query(default=None, max_length=100),
    active_only: bool = True,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*VIEW_ROLES)),
):
    """SDS / kimyasal ürün sicili Excel."""
    stmt = select(ChemicalProduct).order_by(ChemicalProduct.product_name.asc())
    company_ids = company_ids_for_query(db, user, company_id)
    if company_ids == []:
        rows = []
    else:
        if company_ids is not None:
            stmt = stmt.where(ChemicalProduct.company_id.in_(company_ids))
        if active_only:
            stmt = stmt.where(ChemicalProduct.is_active.is_(True))
        if q:
            pattern = f"%{q.strip()}%"
            stmt = stmt.where(
                or_(
                    ChemicalProduct.product_name.ilike(pattern),
                    ChemicalProduct.cas_number.ilike(pattern),
                )
            )
        rows = list(db.scalars(stmt.limit(2000)).all())

    companies = {
        c.id: c.name
        for c in db.scalars(
            select(Company).where(Company.id.in_({r.company_id for r in rows} or {-1}))
        ).all()
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "SDS Sicili"
    headers = [
        "Firma",
        "Ürün",
        "CAS",
        "SDS Dosyası",
        "Doküman ID",
        "Sonraki Gözden Geçirme",
        "Durum",
        "GHS Sayısı",
        "Notlar",
        "Aktif",
    ]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="0D6EFD")
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
    status_tr = {"overdue": "Gecikmiş", "due_soon": "Yaklaşıyor", "ok": "Güncel", "unset": "Tarih yok"}
    for r in rows:
        ghs = parse_checklist(getattr(r, "ghs_checklist_json", None))
        ws.append(
            [
                companies.get(r.company_id, str(r.company_id)),
                r.product_name,
                r.cas_number or "",
                "Var" if r.has_sds_file else "Yok",
                r.document_id or "",
                r.next_review_date.isoformat() if r.next_review_date else "",
                status_tr.get(_review_status(r.next_review_date), _review_status(r.next_review_date)),
                int(ghs["count"]),
                r.notes or "",
                "Evet" if r.is_active else "Hayır",
            ]
        )
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="sds-sicili-{stamp}.xlsx"'},
    )


@router.get("/due-summary", response_model=SdsDueSummary)
def due_summary(
    company_id: int | None = None,
    days: int = Query(30, ge=1, le=365),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*VIEW_ROLES)),
):
    company_ids = company_ids_for_query(db, user, company_id)
    if company_ids == []:
        return SdsDueSummary(
            total=0, with_sds=0, missing_sds=0, due_soon=0, overdue=0, with_ghs_label=0
        )

    q = select(ChemicalProduct).where(ChemicalProduct.is_active.is_(True))
    if company_ids is not None:
        q = q.where(ChemicalProduct.company_id.in_(company_ids))
    rows = list(db.scalars(q).all())
    today = date.today()
    soon = today + timedelta(days=days)
    with_sds = sum(1 for r in rows if r.has_sds_file)
    with_ghs = sum(1 for r in rows if parse_checklist(getattr(r, "ghs_checklist_json", None))["count"] > 0)
    return SdsDueSummary(
        total=len(rows),
        with_sds=with_sds,
        missing_sds=len(rows) - with_sds,
        due_soon=sum(
            1
            for r in rows
            if r.next_review_date and today <= r.next_review_date <= soon
        ),
        overdue=sum(1 for r in rows if r.next_review_date and r.next_review_date < today),
        with_ghs_label=with_ghs,
    )


@router.get("", response_model=list[ChemicalProductResponse])
def list_products(
    company_id: int | None = None,
    q: str | None = Query(default=None, max_length=100),
    active_only: bool = True,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*VIEW_ROLES)),
):
    stmt = select(ChemicalProduct).order_by(ChemicalProduct.product_name.asc())
    company_ids = company_ids_for_query(db, user, company_id)
    if company_ids == []:
        return []
    if company_ids is not None:
        stmt = stmt.where(ChemicalProduct.company_id.in_(company_ids))
    if active_only:
        stmt = stmt.where(ChemicalProduct.is_active.is_(True))
    if q:
        pattern = f"%{q.strip()}%"
        stmt = stmt.where(
            or_(
                ChemicalProduct.product_name.ilike(pattern),
                ChemicalProduct.cas_number.ilike(pattern),
                ChemicalProduct.notes.ilike(pattern),
            )
        )
    return [_to_response(r) for r in db.scalars(stmt).all()]


@router.post("", response_model=ChemicalProductResponse)
def create_product(
    payload: ChemicalProductCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    _ensure_edit(db, user, payload.company_id)
    branch = _validated_branch(db, payload.company_id, payload.branch_id)
    now = datetime.utcnow()
    row = ChemicalProduct(
        company_id=payload.company_id,
        branch_id=branch.id if branch else None,
        product_name=payload.product_name,
        cas_number=payload.cas_number,
        has_sds_file=payload.has_sds_file,
        next_review_date=payload.next_review_date,
        notes=payload.notes,
        is_active=True,
        created_by_id=user.id,
        created_at=now,
        updated_at=now,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return _to_response(row)


@router.patch("/{product_id}", response_model=ChemicalProductResponse)
def update_product(
    product_id: int,
    payload: ChemicalProductUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    row = db.get(ChemicalProduct, product_id)
    if not row:
        raise HTTPException(404, "Kimyasal ürün bulunamadı.")
    _ensure_edit(db, user, row.company_id)
    data = payload.model_dump(exclude_unset=True)
    for key, val in data.items():
        setattr(row, key, val)
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _to_response(row)


@router.post("/{product_id}/ensure-document", response_model=ChemicalProductResponse)
def ensure_sds_document(
    product_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    """Dokümanlar modülüne SDS kaydı oluşturur / bağlar (dosya yükleme için)."""
    row = db.get(ChemicalProduct, product_id)
    if not row:
        raise HTTPException(404, "Kimyasal ürün bulunamadı.")
    _ensure_edit(db, user, row.company_id)

    if row.document_id:
        doc = db.get(DocumentRecord, row.document_id)
        if doc and doc.is_active:
            return _to_response(row)

    title = f"SDS — {row.product_name}"
    if row.cas_number:
        title = f"{title} (CAS {row.cas_number})"
    desc = "PKD/SDS kimyasal ürün sicilinden oluşturuldu."
    if row.notes:
        desc = f"{desc}\n{row.notes}"[:1500]

    doc = DocumentRecord(
        company_id=row.company_id,
        branch_id=row.branch_id,
        category=DocumentCategory.LEGAL,
        title=title[:220],
        file_name=None,
        description=desc,
        valid_from=date.today(),
        valid_until=row.next_review_date,
        version="1.0",
        is_active=True,
        created_by_id=user.id,
    )
    db.add(doc)
    db.flush()
    row.document_id = doc.id
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _to_response(row)


@router.post("/{product_id}/mark-sds-uploaded", response_model=ChemicalProductResponse)
def mark_sds_uploaded(
    product_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    """Doküman dosyası yüklendikten sonra has_sds_file bayrağını günceller."""
    row = db.get(ChemicalProduct, product_id)
    if not row:
        raise HTTPException(404, "Kimyasal ürün bulunamadı.")
    _ensure_edit(db, user, row.company_id)
    if not row.document_id:
        raise HTTPException(400, "Önce SDS doküman kaydı oluşturun.")
    doc = db.get(DocumentRecord, row.document_id)
    if not doc:
        raise HTTPException(404, "Bağlı doküman bulunamadı.")
    has_file = bool(doc.file_name) or ("[stored:" in (doc.description or ""))
    if not has_file:
        raise HTTPException(400, "Dokümana henüz dosya yüklenmemiş.")
    row.has_sds_file = True
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _to_response(row)


@router.get("/{product_id}/ghs-checklist")
def get_ghs_checklist(
    product_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*VIEW_ROLES)),
):
    row = db.get(ChemicalProduct, product_id)
    if not row:
        raise HTTPException(404, "Kimyasal ürün bulunamadı.")
    ensure_company_access(db, user, row.company_id)
    body = parse_checklist(getattr(row, "ghs_checklist_json", None))
    body["product_id"] = row.id
    body["product_name"] = row.product_name
    return body


@router.put("/{product_id}/ghs-checklist", response_model=ChemicalProductResponse)
def put_ghs_checklist(
    product_id: int,
    payload: GhsChecklistUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    """0.9.120 — GHS/CLP tehlike etiketi piktogram checklist stub."""
    row = db.get(ChemicalProduct, product_id)
    if not row:
        raise HTTPException(404, "Kimyasal ürün bulunamadı.")
    _ensure_edit(db, user, row.company_id)
    row.ghs_checklist_json = serialize_selected(payload.selected)
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _to_response(row)


@router.get("/pkd/meta")
def pkd_meta(user: User = Depends(get_current_user)):
    return {
        "engine": PKD_ENGINE,
        "statuses": PKD_STATUS_LABELS,
        "atmosphere_types": PKD_ATMOSPHERE_LABELS,
        "zones": PKD_ZONE_LABELS,
        "ignition_sources": [
            {"code": "hot_work", "label": "Sıcak çalışma / açık alev"},
            {"code": "electrical", "label": "Elektriksel ekipman"},
            {"code": "static", "label": "Statik elektrik"},
            {"code": "mechanical", "label": "Mekanik kıvılcım / sıcak yüzey"},
            {"code": "vehicles", "label": "Araç ve hareketli ekipman"},
            {"code": "lightning", "label": "Yıldırım / atmosferik etki"},
        ],
        "control_measures": [
            {"code": "ventilation", "label": "Yeterli havalandırma"},
            {"code": "ex_equipment", "label": "Uygun Ex ekipman seçimi"},
            {"code": "grounding", "label": "Topraklama ve eşpotansiyel bağlantı"},
            {"code": "hot_work_permit", "label": "Sıcak çalışma izin sistemi"},
            {"code": "gas_detection", "label": "Gaz / toz algılama ve alarm"},
            {"code": "housekeeping", "label": "Toz birikimi ve temizlik kontrolü"},
            {"code": "maintenance", "label": "Periyodik bakım ve muayene"},
            {"code": "training", "label": "Çalışan eğitimi ve bilgilendirme"},
        ],
    }


@router.get("/pkd/export.xlsx")
def export_pkd_xlsx(
    company_id: int | None = None,
    q: str | None = Query(default=None, max_length=100),
    active_only: bool = True,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*VIEW_ROLES)),
):
    """İşyeri/proses bazlı PKD sicilini Excel olarak dışa aktarır."""
    stmt = select(ExplosionProtectionDocument).order_by(
        ExplosionProtectionDocument.area_name.asc(),
        ExplosionProtectionDocument.document_no.asc(),
    )
    company_ids = company_ids_for_query(db, user, company_id)
    if company_ids == []:
        rows = []
    else:
        if company_ids is not None:
            stmt = stmt.where(ExplosionProtectionDocument.company_id.in_(company_ids))
        if active_only:
            stmt = stmt.where(ExplosionProtectionDocument.is_active.is_(True))
        if q:
            pattern = f"%{q.strip()}%"
            stmt = stmt.where(
                or_(
                    ExplosionProtectionDocument.document_no.ilike(pattern),
                    ExplosionProtectionDocument.area_name.ilike(pattern),
                    ExplosionProtectionDocument.process_name.ilike(pattern),
                    ExplosionProtectionDocument.hazardous_materials.ilike(pattern),
                )
            )
        rows = list(db.scalars(stmt.limit(2000)).all())

    companies = {
        c.id: c.name
        for c in db.scalars(
            select(Company).where(Company.id.in_({r.company_id for r in rows} or {-1}))
        ).all()
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "PKD Sicili"
    headers = [
        "Firma",
        "Doküman No",
        "Bölüm / Alan",
        "Proses",
        "Ortam Türü",
        "Zone Sınıfları",
        "Revizyon",
        "Sonraki Gözden Geçirme",
        "Doküman Dosyası",
        "Durum",
        "Sorumlu",
    ]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="0B7D73")
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
    status_tr = {
        "overdue": "Gecikmiş",
        "due_soon": "Yaklaşıyor",
        "ok": "Güncel",
        "unset": "Tarih yok",
        "draft": "Taslak",
        "revision_pending": "Revizyon bekliyor",
        "archived": "Arşiv",
    }
    for row in rows:
        review_status = _pkd_review_status(row)
        zones = ", ".join(PKD_ZONE_LABELS.get(value, value) for value in _json_list(row.zone_classifications_json))
        ws.append(
            [
                companies.get(row.company_id, str(row.company_id)),
                row.document_no,
                row.area_name,
                row.process_name or "",
                PKD_ATMOSPHERE_LABELS.get(row.atmosphere_type, row.atmosphere_type),
                zones,
                row.revision_no or "",
                row.next_review_date.isoformat() if row.next_review_date else "",
                "Var" if row.has_pkd_file else "Yok",
                status_tr.get(review_status, PKD_STATUS_LABELS.get(row.status, row.status)),
                row.responsible_person or "",
            ]
        )
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="pkd-sicili-{stamp}.xlsx"'},
    )


@router.get("/pkd/summary", response_model=PkdSummary)
def pkd_summary(
    company_id: int | None = None,
    days: int = Query(30, ge=1, le=365),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*VIEW_ROLES)),
):
    company_ids = company_ids_for_query(db, user, company_id)
    if company_ids == []:
        return PkdSummary(
            total=0,
            active=0,
            draft=0,
            revision_pending=0,
            with_file=0,
            missing_file=0,
            due_soon=0,
            overdue=0,
        )

    stmt = select(ExplosionProtectionDocument).where(ExplosionProtectionDocument.is_active.is_(True))
    if company_ids is not None:
        stmt = stmt.where(ExplosionProtectionDocument.company_id.in_(company_ids))
    rows = list(db.scalars(stmt).all())
    today = date.today()
    soon = today + timedelta(days=days)
    return PkdSummary(
        total=len(rows),
        active=sum(1 for row in rows if row.status == "active"),
        draft=sum(1 for row in rows if row.status == "draft"),
        revision_pending=sum(1 for row in rows if row.status == "revision_pending"),
        with_file=sum(1 for row in rows if row.has_pkd_file),
        missing_file=sum(1 for row in rows if not row.has_pkd_file),
        due_soon=sum(1 for row in rows if row.next_review_date and today <= row.next_review_date <= soon),
        overdue=sum(1 for row in rows if row.next_review_date and row.next_review_date < today),
    )


@router.get("/pkd", response_model=list[PkdDocumentResponse])
def list_pkd_documents(
    company_id: int | None = None,
    q: str | None = Query(default=None, max_length=100),
    active_only: bool = True,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*VIEW_ROLES)),
):
    stmt = select(ExplosionProtectionDocument).order_by(
        ExplosionProtectionDocument.area_name.asc(),
        ExplosionProtectionDocument.document_no.asc(),
    )
    company_ids = company_ids_for_query(db, user, company_id)
    if company_ids == []:
        return []
    if company_ids is not None:
        stmt = stmt.where(ExplosionProtectionDocument.company_id.in_(company_ids))
    if active_only:
        stmt = stmt.where(ExplosionProtectionDocument.is_active.is_(True))
    if q:
        pattern = f"%{q.strip()}%"
        stmt = stmt.where(
            or_(
                ExplosionProtectionDocument.document_no.ilike(pattern),
                ExplosionProtectionDocument.area_name.ilike(pattern),
                ExplosionProtectionDocument.process_name.ilike(pattern),
                ExplosionProtectionDocument.hazardous_materials.ilike(pattern),
                ExplosionProtectionDocument.notes.ilike(pattern),
            )
        )
    return [_pkd_to_response(row) for row in db.scalars(stmt).all()]


@router.post("/pkd", response_model=PkdDocumentResponse)
def create_pkd_document(
    payload: PkdDocumentCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    _ensure_edit(db, user, payload.company_id)
    branch = _validated_branch(db, payload.company_id, payload.branch_id)
    now = datetime.utcnow()
    row = ExplosionProtectionDocument(
        company_id=payload.company_id,
        branch_id=branch.id if branch else None,
        document_no=payload.document_no,
        area_name=payload.area_name,
        process_name=payload.process_name,
        atmosphere_type=payload.atmosphere_type,
        hazardous_materials=payload.hazardous_materials,
        zone_classifications_json=_serialize_list(payload.zone_classifications),
        ignition_sources_json=_serialize_list(payload.ignition_sources),
        control_measures_json=_serialize_list(payload.control_measures),
        responsible_person=payload.responsible_person,
        prepared_by=payload.prepared_by,
        approved_by=payload.approved_by,
        document_date=payload.document_date,
        revision_no=payload.revision_no,
        next_review_date=payload.next_review_date,
        status=payload.status,
        notes=payload.notes,
        is_active=True,
        created_by_id=user.id,
        created_at=now,
        updated_at=now,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return _pkd_to_response(row)


@router.patch("/pkd/{pkd_id}", response_model=PkdDocumentResponse)
def update_pkd_document(
    pkd_id: int,
    payload: PkdDocumentUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    row = db.get(ExplosionProtectionDocument, pkd_id)
    if not row:
        raise HTTPException(404, "Patlamadan korunma dokümanı bulunamadı.")
    _ensure_edit(db, user, row.company_id)
    data = payload.model_dump(exclude_unset=True)
    if "branch_id" in data:
        branch = _validated_branch(db, row.company_id, data.pop("branch_id"))
        row.branch_id = branch.id if branch else None
    json_fields = {
        "zone_classifications": "zone_classifications_json",
        "ignition_sources": "ignition_sources_json",
        "control_measures": "control_measures_json",
    }
    for key, value in data.items():
        setattr(row, json_fields.get(key, key), _serialize_list(value) if key in json_fields else value)
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _pkd_to_response(row)


@router.post("/pkd/{pkd_id}/ensure-document", response_model=PkdDocumentResponse)
def ensure_pkd_document_record(
    pkd_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    """PKD kaydını ortak Dokümanlar altyapısına bağlar."""
    row = db.get(ExplosionProtectionDocument, pkd_id)
    if not row:
        raise HTTPException(404, "Patlamadan korunma dokümanı bulunamadı.")
    _ensure_edit(db, user, row.company_id)

    if row.document_id:
        doc = db.get(DocumentRecord, row.document_id)
        if doc and doc.is_active:
            return _pkd_to_response(row)

    title = f"PKD — {row.document_no} — {row.area_name}"
    desc = "Patlamadan Korunma Dokümanı sicilinden oluşturuldu."
    if row.process_name:
        desc = f"{desc}\nProses: {row.process_name}"
    if row.notes:
        desc = f"{desc}\n{row.notes}"[:1500]
    doc = DocumentRecord(
        company_id=row.company_id,
        branch_id=row.branch_id,
        category=DocumentCategory.LEGAL,
        title=title[:220],
        file_name=None,
        description=desc,
        valid_from=row.document_date or date.today(),
        valid_until=row.next_review_date,
        version=row.revision_no or "1.0",
        is_active=True,
        created_by_id=user.id,
    )
    db.add(doc)
    db.flush()
    row.document_id = doc.id
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _pkd_to_response(row)


@router.post("/pkd/{pkd_id}/mark-file-uploaded", response_model=PkdDocumentResponse)
def mark_pkd_file_uploaded(
    pkd_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles_or_workplace_operations(*EDIT_ROLES)),
):
    row = db.get(ExplosionProtectionDocument, pkd_id)
    if not row:
        raise HTTPException(404, "Patlamadan korunma dokümanı bulunamadı.")
    _ensure_edit(db, user, row.company_id)
    if not row.document_id:
        raise HTTPException(400, "Önce PKD doküman kaydı oluşturun.")
    doc = db.get(DocumentRecord, row.document_id)
    if not doc:
        raise HTTPException(404, "Bağlı PKD dokümanı bulunamadı.")
    has_file = bool(doc.file_name) or ("[stored:" in (doc.description or ""))
    if not has_file:
        raise HTTPException(400, "PKD dokümanına henüz dosya yüklenmemiş.")
    row.has_pkd_file = True
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _pkd_to_response(row)
