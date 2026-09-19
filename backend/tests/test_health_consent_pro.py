"""Sağlık gözetimi: muayene rızası, aktif atama ve klinik mahremiyet."""
from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture()
def client(tmp_path, monkeypatch):
    url = "sqlite:///:memory:"
    monkeypatch.setenv("DATABASE_URL", url)
    monkeypatch.setenv("ENVIRONMENT", "development")
    monkeypatch.setenv("SECRET_KEY", "test-secret-key-at-least-32-chars-long!!")
    monkeypatch.setattr("app.api.auth.role_requires_mfa", lambda _role: False)

    from app.core.config import settings

    settings.database_url = url
    settings.secret_key = "test-secret-key-at-least-32-chars-long!!"
    settings.environment = "development"
    settings.upload_dir = str(tmp_path / "uploads")

    from sqlalchemy import create_engine
    from sqlalchemy.pool import StaticPool
    from sqlalchemy.orm import sessionmaker

    import app.core.database as dbmod
    import app.models.entities as ent

    engine = create_engine(
        url,
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    dbmod.engine = engine
    dbmod.SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    ent.Base.metadata.create_all(bind=engine)

    from app.core.security import get_password_hash
    from datetime import date
    from app.models.entities import (
        AssignmentStatus,
        Company,
        Employee,
        IsgProfessional,
        OsgbOrganization,
        ProfessionalType,
        User,
        UserRole,
        WorkplaceAssignment,
    )
    from app.main import app

    with dbmod.SessionLocal() as db:
        osgb = OsgbOrganization(name="Onam OSGB", is_active=True)
        db.add(osgb)
        db.flush()
        company = Company(
            name="Onam Firma",
            hazard_class="Tehlikeli",
            osgb_id=osgb.id,
            is_active=True,
        )
        db.add(company)
        db.flush()
        physician = IsgProfessional(
            osgb_id=osgb.id,
            full_name="Hekim Kisi",
            email="onam-hekim@test.com",
            professional_type=ProfessionalType.WORKPLACE_PHYSICIAN,
            certificate_number="HEK-TEST-1",
            is_active=True,
        )
        dsp_professional = IsgProfessional(
            osgb_id=osgb.id,
            full_name="DSP Kisi",
            email="onam-dsp@test.com",
            professional_type=ProfessionalType.OTHER_HEALTH_PERSONNEL,
            certificate_number="DSP-TEST-1",
            is_active=True,
        )
        other_physician = IsgProfessional(
            osgb_id=osgb.id,
            full_name="Baska Hekim",
            email="baska-hekim@test.com",
            professional_type=ProfessionalType.WORKPLACE_PHYSICIAN,
            certificate_number="HEK-TEST-2",
            is_active=True,
        )
        db.add_all([physician, dsp_professional, other_physician])
        db.flush()
        db.add_all([
            WorkplaceAssignment(
                osgb_id=osgb.id,
                company_id=company.id,
                professional_id=physician.id,
                professional_type=ProfessionalType.WORKPLACE_PHYSICIAN,
                start_date=date(2025, 1, 1),
                status=AssignmentStatus.ACTIVE,
            ),
            WorkplaceAssignment(
                osgb_id=osgb.id,
                company_id=company.id,
                professional_id=dsp_professional.id,
                professional_type=ProfessionalType.OTHER_HEALTH_PERSONNEL,
                start_date=date(2025, 1, 1),
                status=AssignmentStatus.ACTIVE,
            ),
            WorkplaceAssignment(
                osgb_id=osgb.id,
                company_id=company.id,
                professional_id=other_physician.id,
                professional_type=ProfessionalType.WORKPLACE_PHYSICIAN,
                start_date=date(2025, 1, 1),
                status=AssignmentStatus.ACTIVE,
            ),
        ])
        db.add_all(
            [
                User(
                    email="onam-hekim@test.com",
                    full_name="Hekim Kisi",
                    hashed_password=get_password_hash("HekimPass123!"),
                    role=UserRole.WORKPLACE_PHYSICIAN,
                    osgb_id=osgb.id,
                    company_id=company.id,
                    is_active=True,
                ),
                User(
                    email="onam-dsp@test.com",
                    full_name="DSP Kisi",
                    hashed_password=get_password_hash("DspPass12345!"),
                    role=UserRole.OTHER_HEALTH_PERSONNEL,
                    osgb_id=osgb.id,
                    company_id=company.id,
                    is_active=True,
                ),
                User(
                    email="onam-osgb@test.com",
                    full_name="Isyeri Yetkilisi",
                    hashed_password=get_password_hash("OsgbPass123!"),
                    role=UserRole.COMPANY_ADMIN,
                    osgb_id=osgb.id,
                    company_id=company.id,
                    is_active=True,
                ),
                User(
                    email="onam-osgb-central@test.com",
                    full_name="OSGB Merkez Yonetici",
                    hashed_password=get_password_hash("CentralPass123!"),
                    role=UserRole.COMPANY_ADMIN,
                    osgb_id=osgb.id,
                    company_id=None,
                    is_active=True,
                ),
                User(
                    email="onam-global@test.com",
                    full_name="EISA Global",
                    hashed_password=get_password_hash("GlobalPass123!"),
                    role=UserRole.GLOBAL_ADMIN,
                    is_active=True,
                ),
                Employee(
                    company_id=company.id,
                    full_name="Personel A",
                    job_title="Teknisyen",
                    department="Uretim",
                    is_active=True,
                ),
            ]
        )
        db.commit()

    return TestClient(app)


def _headers(client: TestClient, email: str, password: str) -> dict:
    r = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _ids() -> tuple[int, int]:
    import app.core.database as dbmod
    from app.models.entities import Employee

    with dbmod.SessionLocal() as db:
        emp = db.query(Employee).first()
        return emp.company_id, emp.id


def _professional_id(email: str) -> int:
    import app.core.database as dbmod
    from app.models.entities import IsgProfessional

    with dbmod.SessionLocal() as db:
        return db.query(IsgProfessional).filter(IsgProfessional.email == email).one().id


def _payload(company_id: int, employee_id: int, **over) -> dict:
    base = {
        "company_id": company_id,
        "employee_id": employee_id,
        "record_type": "periodic_exam",
        "examination_date": "2026-03-10",
        "fitness_status": "fit",
        "summary": "Periyodik muayene bulgulari uygun",
        "informed_consent": True,
    }
    base.update(over)
    return base


def test_create_rejected_without_informed_consent(client):
    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()

    r = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(company_id, employee_id, informed_consent=False),
    )
    assert r.status_code == 400, r.text
    assert "onay" in r.json()["detail"].casefold()


def test_create_with_consent_stamps_time(client):
    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()

    r = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(company_id, employee_id),
    )
    assert r.status_code in (200, 201), r.text
    body = r.json()
    assert body["informed_consent"] is True
    assert body["informed_consent_at"]


def test_duplicate_exam_same_day_rejected(client):
    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()

    first = client.post(
        "/api/v1/health-records", headers=headers, json=_payload(company_id, employee_id)
    )
    assert first.status_code in (200, 201), first.text

    again = client.post(
        "/api/v1/health-records", headers=headers, json=_payload(company_id, employee_id)
    )
    assert again.status_code == 400, again.text
    assert "zaten var" in again.json()["detail"].casefold()

    # Farklı tarih engellenmez
    other_day = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(company_id, employee_id, examination_date="2026-04-11"),
    )
    assert other_day.status_code in (200, 201), other_day.text


def test_periodic_exam_rejects_date_beyond_hazard_ceiling(client):
    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    rejected = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(
            company_id,
            employee_id,
            examination_date="2026-03-10",
            next_examination_date="2030-03-10",
        ),
    )
    assert rejected.status_code == 422, rejected.text
    assert "azami periyodu" in rejected.json()["detail"]


def test_special_policy_employee_defaults_to_six_months(client):
    import app.core.database as dbmod
    from app.models.entities import Employee

    with dbmod.SessionLocal() as db:
        employee = db.query(Employee).first()
        employee.special_status = "Gebe"
        db.commit()

    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    created = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(company_id, employee_id, examination_date="2026-03-10"),
    )
    assert created.status_code in (200, 201), created.text
    assert created.json()["next_examination_date"] == "2026-09-10"


def test_workplace_manager_gets_full_read_and_safe_downloads(client):
    physician = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    restrictions = "Gece vardiyasında çalışamaz; 10 kg üstü yük kaldıramaz"
    secrets = {
        "summary": "KLINIK_OZET_GIZLI",
        "confidential_note": "HEKIM_NOTU_GIZLI",
        "audiometry_result": "ODYO_SONUCU_GIZLI",
        "spirometry_result": "SFT_SONUCU_GIZLI",
        "chest_xray_result": "AKCIGER_SONUCU_GIZLI",
        "suggested_tests": "TETKIK_ONERISI_GIZLI",
        "exposures": "MARUZIYET_GIZLI",
        "follow_up_note": "TAKIP_NOTU_GIZLI",
        "other_biological_test": "BIYOLOJIK_TEST_GIZLI",
    }
    created = client.post(
        "/api/v1/health-records",
        headers=physician,
        json=_payload(
            company_id,
            employee_id,
            examination_date="2026-08-10",
            next_examination_date="2027-08-10",
            fitness_status="conditional",
            restrictions=restrictions,
            audiometry_date="2026-08-10",
            spirometry_date="2026-08-10",
            chest_xray_date="2026-08-10",
            blood_lead_date="2026-08-10",
            blood_lead_value=44,
            blood_lead_ref=30,
            **secrets,
        ),
    )
    assert created.status_code in (200, 201), created.text
    record_id = created.json()["id"]

    uploaded = client.post(
        f"/api/v1/health-records/{record_id}/report",
        headers=physician,
        files={"file": ("muayene-raporu.pdf", b"%PDF-1.4\n", "application/pdf")},
    )
    assert uploaded.status_code == 200, uploaded.text

    headers = _headers(client, "onam-osgb@test.com", "OsgbPass123!")

    meta = client.get("/api/v1/health-records/meta", headers=headers)
    assert meta.status_code == 200, meta.text

    listed = client.get(f"/api/v1/health-records?company_id={company_id}", headers=headers)
    assert listed.status_code == 200, listed.text
    row = next(item for item in listed.json() if item["id"] == record_id)
    assert row["employee_name"] == "Personel A"
    assert row["fitness_status"] == "conditional"
    assert row["restrictions"] == restrictions
    assert row["physician_name"] == "Hekim Kisi"
    assert row["informed_consent"] is True
    assert row["informed_consent_at"] is not None
    assert row["has_report"] is True
    assert row["physician_professional_id"] is not None
    for field, value in secrets.items():
        assert row[field] == value, field
    assert row["audiometry_date"] == "2026-08-10"
    assert row["spirometry_date"] == "2026-08-10"
    assert row["chest_xray_date"] == "2026-08-10"
    assert row["blood_lead_date"] == "2026-08-10"
    assert row["blood_lead_value"] == 44
    assert row["blood_lead_ref"] == 30
    assert row["blood_lead_eval"] == "yuksek"
    assert row["report_file_name"] == "muayene-raporu.pdf"

    summary = client.get(
        f"/api/v1/health-records/summary?company_id={company_id}", headers=headers
    )
    assert summary.status_code == 200, summary.text
    assert summary.json()["conditional"] == 1
    for field in ("with_audiometry", "with_spirometry", "with_chest_xray", "with_blood_lead", "lead_high"):
        assert summary.json()[field] == 1

    fitness = client.get(
        f"/api/v1/health-records/{record_id}/fitness.html", headers=headers
    )
    assert fitness.status_code == 200, fitness.text
    assert restrictions in fitness.text
    for secret in secrets.values():
        assert secret not in fitness.text

    full_page = client.get(
        f"/api/v1/health-records/{record_id}/form.html", headers=headers
    )
    assert full_page.status_code == 200, full_page.text
    assert "Çalışan Sağlık Bilgileri — Salt Okunur" in full_page.text
    for secret in secrets.values():
        assert secret in full_page.text

    report = client.get(
        f"/api/v1/health-records/{record_id}/report", headers=headers
    )
    assert report.status_code == 200, report.text
    assert report.headers["content-type"].startswith("application/pdf")
    assert report.content.startswith(b"%PDF")

    exported = client.get(
        f"/api/v1/health-records/export.xlsx?company_id={company_id}", headers=headers
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(BytesIO(exported.content), data_only=False)
    assert workbook.sheetnames == ["Sağlık Gözetimi"]
    values = [cell.value for row_cells in workbook.active.iter_rows() for cell in row_cells]
    exported_text = "\n".join(str(value) for value in values if value is not None)
    assert "Personel A" in exported_text
    assert restrictions in exported_text
    assert "Kan Kurşun" in exported_text
    assert "44" in exported_text
    for secret in secrets.values():
        assert secret in exported_text

    create_attempt = client.post(
        "/api/v1/health-records", headers=headers, json=_payload(company_id, employee_id)
    )
    assert create_attempt.status_code == 403, create_attempt.text
    assert client.patch(
        f"/api/v1/health-records/{record_id}",
        headers=headers,
        json={"restrictions": "Yetkisiz değişiklik"},
    ).status_code == 403
    assert client.delete(
        f"/api/v1/health-records/{record_id}?reason=Yetkisiz+silme",
        headers=headers,
    ).status_code == 403
    assert client.post(
        f"/api/v1/health-records/{record_id}/report",
        headers=headers,
        files={"file": ("rapor.pdf", b"%PDF-1.4\n", "application/pdf")},
    ).status_code == 403


def test_osgb_central_admin_cannot_reach_health_records(client):
    """OSGB merkez yöneticisi çalışanların tıbbi kaydını göremez/açamaz."""
    headers = _headers(client, "onam-osgb-central@test.com", "CentralPass123!")
    company_id, employee_id = _ids()

    assert client.get(
        f"/api/v1/health-records?company_id={company_id}", headers=headers
    ).status_code == 403
    assert client.get("/api/v1/health-records/meta", headers=headers).status_code == 403
    assert client.post(
        "/api/v1/health-records", headers=headers, json=_payload(company_id, employee_id)
    ).status_code == 403


def test_global_admin_cannot_reach_clinical_health_records(client):
    headers = _headers(client, "onam-global@test.com", "GlobalPass123!")
    company_id, employee_id = _ids()
    listed = client.get(f"/api/v1/health-records?company_id={company_id}", headers=headers)
    assert listed.status_code == 403, listed.text
    created = client.post(
        "/api/v1/health-records", headers=headers, json=_payload(company_id, employee_id)
    )
    assert created.status_code == 403, created.text


def test_restrictions_encrypted_at_rest_and_hidden_from_dsp(client, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "health_field_encryption_enabled", True)
    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    limit_text = "Agir kaldirma ve gece vardiyasi yapmamali"

    created = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(company_id, employee_id, restrictions=limit_text),
    )
    assert created.status_code in (200, 201), created.text
    record_id = created.json()["id"]
    assert created.json()["restrictions"] == limit_text

    import app.core.database as dbmod
    from app.models.entities import HealthRecord

    with dbmod.SessionLocal() as db:
        stored = db.get(HealthRecord, record_id).restrictions
    assert stored.startswith("enc:v1:")
    assert limit_text not in stored

    dsp = _headers(client, "onam-dsp@test.com", "DspPass12345!")
    rows = client.get(f"/api/v1/health-records?company_id={company_id}", headers=dsp)
    assert rows.status_code == 200, rows.text
    row = next(x for x in rows.json() if x["id"] == record_id)
    assert row["restrictions"] is None
    assert row["fitness_status"] is None
    assert row["smart_summary"] is None
    assert row["tetkik_summary"] is None


def test_dsp_cannot_search_or_filter_clinical_decisions(client):
    physician = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    created = client.post(
        "/api/v1/health-records",
        headers=physician,
        json=_payload(
            company_id,
            employee_id,
            examination_date="2026-08-10",
            summary="KLINIK_OZET_GIZLI",
        ),
    )
    assert created.status_code in (200, 201), created.text

    dsp = _headers(client, "onam-dsp@test.com", "DspPass12345!")
    clinical_search = client.get(
        f"/api/v1/health-records?company_id={company_id}&q=KLINIK_OZET_GIZLI",
        headers=dsp,
    )
    assert clinical_search.status_code == 200, clinical_search.text
    assert clinical_search.json() == []

    employee_search = client.get(
        f"/api/v1/health-records?company_id={company_id}&q=Personel",
        headers=dsp,
    )
    assert employee_search.status_code == 200, employee_search.text
    assert len(employee_search.json()) == 1

    status_filter = client.get(
        f"/api/v1/health-records?company_id={company_id}&fitness_status=fit",
        headers=dsp,
    )
    assert status_filter.status_code == 403, status_filter.text


def test_physician_identity_is_locked_to_own_active_assignment(client):
    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    other_id = _professional_id("baska-hekim@test.com")
    rejected = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(
            company_id,
            employee_id,
            examination_date="2026-05-10",
            physician_professional_id=other_id,
        ),
    )
    assert rejected.status_code == 403, rejected.text

    created = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(company_id, employee_id, examination_date="2026-05-11"),
    )
    assert created.status_code in (200, 201), created.text
    assert created.json()["physician_professional_id"] == _professional_id("onam-hekim@test.com")
    assert created.json()["physician_name"] == "Hekim Kisi"


def test_dsp_field_matrix_blocks_clinical_decisions(client):
    dsp = _headers(client, "onam-dsp@test.com", "DspPass12345!")
    company_id, employee_id = _ids()
    physician_id = _professional_id("onam-hekim@test.com")

    forbidden = client.post(
        "/api/v1/health-records",
        headers=dsp,
        json=_payload(
            company_id,
            employee_id,
            examination_date="2026-06-10",
            physician_professional_id=physician_id,
        ),
    )
    assert forbidden.status_code == 403, forbidden.text
    assert "klinik" in forbidden.json()["detail"].casefold()

    allowed = client.post(
        "/api/v1/health-records",
        headers=dsp,
        json={
            "company_id": company_id,
            "employee_id": employee_id,
            "record_type": "periodic_exam",
            "examination_date": "2026-06-11",
            "physician_professional_id": physician_id,
            "informed_consent": True,
            "audiometry_date": "2026-06-11",
            "blood_lead_date": "2026-06-11",
            "blood_lead_value": 12.5,
            "blood_lead_unit": "µg/dL",
        },
    )
    assert allowed.status_code in (200, 201), allowed.text
    record_id = allowed.json()["id"]
    # DSP klinik uygunluk kararını göremez; yalnız teknik kaydı oluşturur.
    assert allowed.json()["fitness_status"] is None

    clinical_patch = client.patch(
        f"/api/v1/health-records/{record_id}",
        headers=dsp,
        json={"summary": "DSP klinik yorum yazamaz"},
    )
    assert clinical_patch.status_code == 403, clinical_patch.text
    deleted = client.delete(
        f"/api/v1/health-records/{record_id}?reason=Yanlis kayit",
        headers=dsp,
    )
    assert deleted.status_code == 403, deleted.text
    clinical_form = client.get(f"/api/v1/health-records/{record_id}/form.html", headers=dsp)
    assert clinical_form.status_code == 403, clinical_form.text


def test_clinical_and_employer_documents_are_separated(client):
    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    created = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(
            company_id,
            employee_id,
            examination_date="2026-07-10",
            summary="KLINIK_OZET_GIZLI",
            confidential_note="HEKIM_NOTU_GIZLI",
            restrictions="Ağır kaldırma yapamaz",
            audiometry_result="ODYO_SONUCU_GIZLI",
        ),
    )
    assert created.status_code in (200, 201), created.text
    record_id = created.json()["id"]

    clinical = client.get(f"/api/v1/health-records/{record_id}/form.html", headers=headers)
    assert clinical.status_code == 200, clinical.text
    assert "Gizli Klinik Sağlık Dosyası" in clinical.text
    assert "İşe Giriş / Periyodik Muayene Formu" in clinical.text
    assert "KLINIK_OZET_GIZLI" in clinical.text
    assert "İşveren / Vekili" not in clinical.text

    fitness = client.get(f"/api/v1/health-records/{record_id}/fitness.html", headers=headers)
    assert fitness.status_code == 200, fitness.text
    assert "İşe Uygunluk ve Çalışma Kısıtları Belgesi" in fitness.text
    assert "Ağır kaldırma yapamaz" in fitness.text
    assert "KLINIK_OZET_GIZLI" not in fitness.text
    assert "HEKIM_NOTU_GIZLI" not in fitness.text
    assert "ODYO_SONUCU_GIZLI" not in fitness.text


def test_revision_and_access_logs_are_hash_chained_and_append_only(client):
    from sqlalchemy import text
    from sqlalchemy.exc import DatabaseError
    import app.core.database as dbmod
    from app.models.entities import HealthAccessLog, HealthRecordRevision
    from app.services.health_audit import verify_access_chain, verify_revision_chain

    headers = _headers(client, "onam-hekim@test.com", "HekimPass123!")
    company_id, employee_id = _ids()
    created = client.post(
        "/api/v1/health-records",
        headers=headers,
        json=_payload(company_id, employee_id, examination_date="2026-07-11"),
    )
    assert created.status_code in (200, 201), created.text
    record_id = created.json()["id"]
    updated = client.patch(
        f"/api/v1/health-records/{record_id}?change_reason=Kontrol+sonucu+guncellendi",
        headers=headers,
        json={"summary": "Kontrol sonucu güncellendi"},
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["version"] == 2
    listed = client.get(f"/api/v1/health-records?company_id={company_id}", headers=headers)
    assert listed.status_code == 200, listed.text

    with dbmod.SessionLocal() as db:
        revisions = list(
            db.query(HealthRecordRevision)
            .filter(HealthRecordRevision.record_id == record_id)
            .order_by(HealthRecordRevision.version)
        )
        accesses = list(
            db.query(HealthAccessLog)
            .filter(HealthAccessLog.company_id == company_id)
            .order_by(HealthAccessLog.id)
        )
        assert [r.action for r in revisions] == ["create", "update"]
        assert verify_revision_chain(revisions) is True
        assert verify_access_chain(accesses) is True
        with pytest.raises(DatabaseError):
            db.execute(
                text("UPDATE health_record_revisions SET action='tampered' WHERE id=:id"),
                {"id": revisions[0].id},
            )
            db.commit()
        db.rollback()
        with pytest.raises(DatabaseError):
            db.execute(
                text("DELETE FROM health_access_logs WHERE id=:id"),
                {"id": accesses[0].id},
            )
            db.commit()
