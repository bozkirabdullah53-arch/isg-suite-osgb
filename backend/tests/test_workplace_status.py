"""İşyeri Durum Merkezi: tenant izolasyonu, mahremiyet ve rapor smoke testleri."""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture()
def client(tmp_path, monkeypatch):
    db_file = tmp_path / "workplace-status.db"
    url = f"sqlite:///{db_file.as_posix()}"
    monkeypatch.setenv("DATABASE_URL", url)
    monkeypatch.setenv("ENVIRONMENT", "development")
    monkeypatch.setenv("SECRET_KEY", "test-secret-key-at-least-32-chars-long!!")
    monkeypatch.setattr("app.api.auth.role_requires_mfa", lambda _role: False)

    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    import app.core.database as dbmod
    import app.models.entities as ent
    from app.core.config import settings

    settings.database_url = url
    settings.secret_key = "test-secret-key-at-least-32-chars-long!!"
    settings.environment = "development"
    engine = create_engine(url, connect_args={"check_same_thread": False})
    dbmod.engine = engine
    dbmod.SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    ent.Base.metadata.create_all(bind=engine)

    from app.main import app

    return TestClient(app)


def _seed() -> dict:
    from app.core.database import SessionLocal
    from app.core.security import get_password_hash
    from app.models.entities import (
        AssignmentStatus,
        Company,
        Employee,
        HealthFitnessStatus,
        HealthRecord,
        HealthRecordType,
        IsgProfessional,
        OsgbOrganization,
        ProfessionalType,
        User,
        UserRole,
        WorkplaceAssignment,
    )

    password = "TestPass123!"
    with SessionLocal() as db:
        osgb_1 = OsgbOrganization(name="Birinci OSGB", authorization_number="OSGB-1", is_active=True)
        osgb_2 = OsgbOrganization(name="İkinci OSGB", authorization_number="OSGB-2", is_active=True)
        db.add_all([osgb_1, osgb_2])
        db.flush()
        company_1 = Company(name="Yetkili İşyeri", osgb_id=osgb_1.id, is_active=True, hazard_class="Tehlikeli")
        company_2 = Company(name="Yabancı İşyeri", osgb_id=osgb_2.id, is_active=True, hazard_class="Tehlikeli")
        db.add_all([company_1, company_2])
        db.flush()

        admin = User(
            email="admin@birinci.example.com", full_name="Birinci OSGB Admin", hashed_password=get_password_hash(password),
            role=UserRole.COMPANY_ADMIN, osgb_id=osgb_1.id, is_active=True,
        )
        employer = User(
            email="employer@birinci.example.com", full_name="İşveren", hashed_password=get_password_hash(password),
            role=UserRole.READ_ONLY, company_id=company_1.id, osgb_id=osgb_1.id, is_active=True,
        )
        workplace_manager = User(
            email="ik@birinci.example.com", full_name="İK Yetkilisi", hashed_password=get_password_hash(password),
            role=UserRole.COMPANY_ADMIN, company_id=company_1.id, osgb_id=osgb_1.id, is_active=True,
        )
        specialist = User(
            email="uzman@birinci.example.com", full_name="Atanmış Uzman", hashed_password=get_password_hash(password),
            role=UserRole.SAFETY_SPECIALIST, osgb_id=osgb_1.id, is_active=True,
        )
        db.add_all([admin, employer, workplace_manager, specialist])
        db.flush()
        professional = IsgProfessional(
            osgb_id=osgb_1.id,
            full_name="Atanmış Uzman",
            email="uzman@birinci.example.com",
            professional_type=ProfessionalType.SAFETY_SPECIALIST,
            certificate_class="A",
            certificate_number="A-100",
            is_active=True,
        )
        db.add(professional)
        db.flush()
        db.add(
            WorkplaceAssignment(
                osgb_id=osgb_1.id,
                company_id=company_1.id,
                professional_id=professional.id,
                professional_type=ProfessionalType.SAFETY_SPECIALIST,
                start_date=date.today(),
                status=AssignmentStatus.ACTIVE,
            )
        )
        employee = Employee(company_id=company_1.id, full_name="Gizli Çalışan", is_active=True)
        db.add(employee)
        db.flush()
        db.add(
            HealthRecord(
                company_id=company_1.id,
                employee_id=employee.id,
                record_type=HealthRecordType.PERIODIC_EXAM,
                examination_date=date.today() - timedelta(days=300),
                next_examination_date=date.today() - timedelta(days=1),
                fitness_status=HealthFitnessStatus.CONDITIONAL,
                summary="GİZLİ_TANI_METNİ",
                confidential_note="GİZLİ_HEKİM_NOTU",
                restrictions="GİZLİ_KISITLAMA",
                created_by_id=admin.id,
            )
        )
        db.commit()
        return {
            "password": password,
            "company_1": company_1.id,
            "company_2": company_2.id,
            "employee_id": employee.id,
            "admin_id": admin.id,
            "users": [admin.email, employer.email, workplace_manager.email, specialist.email],
        }


def _token(client: TestClient, email: str, password: str) -> str:
    response = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


def test_status_isolated_for_admin_employer_and_assigned_specialist(client):
    seed = _seed()
    for email in seed["users"]:
        headers = {"Authorization": f"Bearer {_token(client, email, seed['password'])}"}
        own = client.get(f"/api/v1/companies/{seed['company_1']}/status", headers=headers)
        assert own.status_code == 200, (email, own.text)
        assert own.headers["cache-control"] == "no-store"
        assert own.json()["company"]["id"] == seed["company_1"]
        if email != seed["users"][0]:
            assert "finance" not in own.json()
            assert "contracts" not in own.json()

        foreign = client.get(f"/api/v1/companies/{seed['company_2']}/status", headers=headers)
        assert foreign.status_code == 403, (email, foreign.text)


def test_status_never_exposes_sensitive_medical_fields_or_ibys_ready_claim(client):
    seed = _seed()
    headers = {"Authorization": f"Bearer {_token(client, seed['users'][0], seed['password'])}"}
    response = client.get(f"/api/v1/companies/{seed['company_1']}/status", headers=headers)
    assert response.status_code == 200, response.text
    payload = response.json()
    serialized = response.text
    assert "GİZLİ_TANI_METNİ" not in serialized
    assert "GİZLİ_HEKİM_NOTU" not in serialized
    assert "GİZLİ_KISITLAMA" not in serialized
    assert "unfit" not in payload["health"]
    assert payload["status_center"]["privacy"] == {
        "medical_data_mode": "aggregate_only",
        "sensitive_fields_exposed": False,
    }
    ibys = payload["status_center"]["ibys_validation"]
    assert ibys["officially_verified"] is False
    assert ibys["readiness_claim"] is False
    assert ibys["status"] == "pending_official_validation"


def test_empty_workplace_does_not_report_empty_modules_as_completed(client):
    seed = _seed()
    from app.core.database import SessionLocal
    from app.models.entities import Company

    with SessionLocal() as db:
        source = db.get(Company, seed["company_1"])
        empty = Company(
            name="Kayıtsız İşyeri",
            osgb_id=source.osgb_id,
            hazard_class="Tehlikeli",
            is_active=True,
        )
        db.add(empty)
        db.commit()
        empty_id = empty.id

    headers = {"Authorization": f"Bearer {_token(client, seed['users'][0], seed['password'])}"}
    response = client.get(f"/api/v1/companies/{empty_id}/status", headers=headers)
    assert response.status_code == 200, response.text

    items = {row["code"]: row for row in response.json()["status_center"]["items"]}
    assert items["health_examinations"]["status"] == "informational"
    assert items["health_examinations"]["status_label"] == "Bilgi"
    assert items["capa"]["status"] == "informational"
    assert items["capa"]["status_label"] == "Bilgi"
    assert items["capa"]["count"] == 0
    assert "Tamamlandı" not in {
        items["health_examinations"]["status_label"],
        items["capa"]["status_label"],
    }


def test_status_combines_workplace_deadlines_without_cross_company_data(client):
    seed = _seed()
    from app.core.database import SessionLocal
    from app.models.entities import (
        ChemicalProduct,
        Company,
        DrillRecord,
        OhsCommitteeMeeting,
        PeriodicControl,
        PpeAssignment,
        TrainingSession,
        TrainingStatus,
        WorkplaceMeasurement,
    )
    from app.models.remote_training import RemoteTrainingAssignment

    today = date.today()
    with SessionLocal() as db:
        RemoteTrainingAssignment.__table__.create(bind=db.get_bind(), checkfirst=True)
        company = db.get(Company, seed["company_1"])
        company.risk_assessment_date = today - timedelta(days=4 * 365)
        db.add_all([
            PpeAssignment(
                company_id=seed["company_1"], employee_id=seed["employee_id"],
                delivery_date=today - timedelta(days=300), category="Baş", item_type="Baret",
                renewal_date=today + timedelta(days=7), created_by_id=seed["admin_id"],
            ),
            TrainingSession(
                company_id=seed["company_1"], title="Temel İSG Yenileme",
                start_date=today - timedelta(days=365), next_training_date=today - timedelta(days=1),
                duration_hours=8, renewal_years=1, hazard_class="Tehlikeli",
                instructor_name="Test Uzmanı", status=TrainingStatus.COMPLETED,
                created_by_id=seed["admin_id"],
            ),
            DrillRecord(
                company_id=seed["company_1"], drill_type="Yangın",
                drill_date=today + timedelta(days=5), status="planlandi",
                scenario="Yangın tahliye senaryosu", created_by_id=seed["admin_id"],
            ),
            PeriodicControl(
                company_id=seed["company_1"], category="kaldirma",
                equipment_name="Yük asansörü", next_due_date=today - timedelta(days=2),
                created_by_id=seed["admin_id"],
            ),
            WorkplaceMeasurement(
                company_id=seed["company_1"], measurement_type="Gürültü",
                location="Üretim", measured_at=today - timedelta(days=300),
                next_due_date=today + timedelta(days=10), created_by_id=seed["admin_id"],
            ),
            ChemicalProduct(
                company_id=seed["company_1"], product_name="Test Kimyasalı",
                has_sds_file=True, next_review_date=today + timedelta(days=8),
                created_by_id=seed["admin_id"],
            ),
            OhsCommitteeMeeting(
                company_id=seed["company_1"], meeting_date=today - timedelta(days=20),
                next_meeting_date=today + timedelta(days=9), created_by_id=seed["admin_id"],
            ),
            RemoteTrainingAssignment(
                company_id=seed["company_1"], program_id=999,
                employee_id=seed["employee_id"], employee_name_snapshot="Gizli Çalışan",
                status="in_progress", due_date=today + timedelta(days=6),
                assigned_by_id=seed["admin_id"],
            ),
        ])
        db.commit()

    headers = {"Authorization": f"Bearer {_token(client, seed['users'][2], seed['password'])}"}
    response = client.get(f"/api/v1/companies/{seed['company_1']}/status", headers=headers)
    assert response.status_code == 200, response.text
    payload = response.json()
    center = payload["status_center"]

    assert "finance" not in payload
    assert "contracts" not in payload
    assert center["schema_version"] == "1.1"
    assert center["deadline_summary"]["overdue"] >= 2
    assert center["deadline_summary"]["due_soon"] >= 5
    assert {row["source"] for row in center["deadlines"]}.issuperset({
        "Risk Değerlendirmesi",
        "Eğitim Yenileme",
        "Uzaktan Eğitim",
        "KKD Değişim",
        "Tatbikat",
        "Periyodik Kontrol",
        "SDS / PKD",
        "Ortam Ölçümü",
        "İSG Kurulu",
    })
    assert {row["code"] for row in center["items"]}.issuperset({
        "risk_assessment_validity",
        "training",
        "ppe",
        "drills",
        "periodic_controls",
        "sds",
        "workplace_measurements",
        "ohs_committee",
    })


def test_risk_assessment_status_is_separate_from_overdue_risk_terms(client):
    seed = _seed()
    from app.core.database import SessionLocal
    from app.models.entities import Company, Hazard, HazardCategory, RiskAssessment

    today = date.today()
    with SessionLocal() as db:
        company = db.get(Company, seed["company_1"])
        company.risk_assessment_date = today - timedelta(days=10)
        category = HazardCategory(name="Durum testi", sort_order=1)
        db.add(category)
        db.flush()
        hazard = Hazard(category_id=category.id, code="H-STATUS-01", name="Test tehlikesi")
        db.add(hazard)
        db.flush()
        db.add(
            RiskAssessment(
                risk_code="R-STATUS-01",
                company_id=seed["company_1"],
                hazard_id=hazard.id,
                activity="Test faaliyeti",
                risk_definition="Termin tarihi geçmiş açık risk",
                probability=2,
                severity=2,
                risk_score=4,
                risk_level="Düşük",
                status="Açık",
                term_date=today - timedelta(days=1),
                created_by_id=seed["admin_id"],
            )
        )
        db.commit()

    headers = {"Authorization": f"Bearer {_token(client, seed['users'][0], seed['password'])}"}
    response = client.get(f"/api/v1/companies/{seed['company_1']}/status", headers=headers)
    assert response.status_code == 200, response.text
    items = {row["code"]: row for row in response.json()["status_center"]["items"]}
    assert items["risk_assessment"]["status"] == "completed"
    assert items["risk_assessment"]["status_label"] == "Kayıtlı"
    assert items["risk_assessment"]["count"] == 1
    assert "1 geçmiş termin tarihi" in items["risk_assessment"]["detail"]
    assert items["risk_assessment_validity"]["status"] == "completed"
    assert items["risk_assessment_validity"]["status_label"] == "Geçerli"

    obligations = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/obligations",
        params={"category": "risk", "status": "overdue"},
        headers=headers,
    )
    assert obligations.status_code == 200, obligations.text
    overdue_risks = [row for row in obligations.json()["items"] if row["source"] == "Risk"]
    assert len(overdue_risks) == 1
    assert overdue_risks[0]["status"] == "overdue"


def test_obligations_are_paginated_filterable_and_keep_renewals_separate(client):
    seed = _seed()
    from app.core.database import SessionLocal
    from app.models.entities import Branch, PeriodicControl, TrainingSession, TrainingStatus

    today = date.today()
    with SessionLocal() as db:
        own_branch = Branch(company_id=seed["company_1"], name="Üretim Şubesi", is_active=True)
        foreign_branch = Branch(company_id=seed["company_2"], name="Yabancı Şube", is_active=True)
        db.add_all([own_branch, foreign_branch])
        db.flush()

        due_dates = [
            today - timedelta(days=2),
            today + timedelta(days=7),
            today + timedelta(days=8),
            *[today + timedelta(days=60 + index) for index in range(104)],
        ]
        db.add_all([
            PeriodicControl(
                company_id=seed["company_1"],
                category="kaldirma",
                equipment_name=f"Ekipman {index:03d}",
                next_due_date=due,
                created_by_id=seed["admin_id"],
            )
            for index, due in enumerate(due_dates)
        ])
        training = TrainingSession(
            company_id=seed["company_1"],
            branch_id=own_branch.id,
            title="Şube Temel İSG Eğitimi",
            start_date=today - timedelta(days=90),
            end_date=today - timedelta(days=89),
            next_training_date=today + timedelta(days=4),
            duration_hours=8,
            renewal_years=1,
            hazard_class="Tehlikeli",
            instructor_name="Test Uzmanı",
            status=TrainingStatus.COMPLETED,
            created_by_id=seed["admin_id"],
        )
        db.add(training)
        db.commit()
        own_branch_id = own_branch.id
        foreign_branch_id = foreign_branch.id
        training_id = training.id

    headers = {"Authorization": f"Bearer {_token(client, seed['users'][2], seed['password'])}"}
    response = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/obligations",
        params={"category": "periodic_control", "page": 1, "page_size": 50},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["pagination"] == {
        "page": 1,
        "page_size": 50,
        "total": 107,
        "total_pages": 3,
        "has_previous": False,
        "has_next": True,
    }
    assert len(payload["items"]) == 50
    assert payload["items"][0]["status"] == "overdue"
    assert payload["items"][0]["target"]["entity_type"] == "periodic_control"
    assert payload["summary"]["overdue"] == 1
    assert payload["summary"]["very_soon"] == 1
    assert payload["summary"]["approaching"] == 1
    assert payload["summary"]["scheduled"] == 104

    overdue = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/obligations",
        params={"category": "periodic_control", "status": "overdue"},
        headers=headers,
    )
    assert overdue.status_code == 200, overdue.text
    assert overdue.json()["pagination"]["total"] == 1

    training_response = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/obligations",
        params={"category": "training", "branch_id": own_branch_id},
        headers=headers,
    )
    assert training_response.status_code == 200, training_response.text
    training_items = training_response.json()["items"]
    assert {row["status"] for row in training_items} == {"very_soon", "completed"}
    assert {row["target"]["record_id"] for row in training_items} == {training_id}
    assert len({row["key"] for row in training_items}) == 2

    foreign_branch = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/obligations",
        params={"branch_id": foreign_branch_id},
        headers=headers,
    )
    assert foreign_branch.status_code == 404

    foreign_company = client.get(
        f"/api/v1/companies/{seed['company_2']}/status/obligations",
        headers=headers,
    )
    assert foreign_company.status_code == 403


def test_one_click_pdf_and_excel_reports_are_valid_and_scoped(client):
    seed = _seed()
    headers = {"Authorization": f"Bearer {_token(client, seed['users'][0], seed['password'])}"}
    pdf = client.get(f"/api/v1/companies/{seed['company_1']}/status/report.pdf", headers=headers)
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert pdf.headers["cache-control"] == "no-store"

    excel = client.get(f"/api/v1/companies/{seed['company_1']}/status/report.xlsx", headers=headers)
    assert excel.status_code == 200, excel.text
    assert excel.content.startswith(b"PK")
    assert "spreadsheetml" in excel.headers["content-type"]
    assert excel.headers["cache-control"] == "no-store"

    foreign = client.get(f"/api/v1/companies/{seed['company_2']}/status/report.pdf", headers=headers)
    assert foreign.status_code == 403


def test_full_company_file_keeps_commercial_scope_and_adds_detail_sheets(client):
    seed = _seed()
    from app.core.database import SessionLocal
    from app.models.entities import Branch, FinanceTransaction, ServiceContract

    with SessionLocal() as db:
        db.add(Branch(
            company_id=seed["company_1"],
            name="Üretim Şubesi",
            sgk_registry_no="SGK-SUBE-1",
            city="Balıkesir",
            address="Organize Sanayi Bölgesi",
            is_active=True,
        ))
        db.add(ServiceContract(
            osgb_id=1,
            company_id=seed["company_1"],
            contract_number="SOZ-2026-01",
            start_date=date.today(),
            end_date=date.today() + timedelta(days=180),
            monthly_fee=12500,
            status="active",
        ))
        db.add_all([
            FinanceTransaction(
                osgb_id=1,
                company_id=seed["company_1"],
                transaction_type="income",
                category="contract",
                amount=12500,
                transaction_date=date.today(),
                due_date=date.today() - timedelta(days=2),
                status="pending",
                description="Ocak sözleşme tahakkuku",
            ),
            FinanceTransaction(
                osgb_id=1,
                company_id=seed["company_1"],
                transaction_type="income",
                category="service",
                amount=10000,
                transaction_date=date.today(),
                status="paid",
                description="Tahsilat",
            ),
        ])
        db.commit()

    admin_headers = {"Authorization": f"Bearer {_token(client, seed['users'][0], seed['password'])}"}
    full_pdf = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/full-report.pdf",
        headers=admin_headers,
    )
    assert full_pdf.status_code == 200, full_pdf.text
    assert full_pdf.content.startswith(b"%PDF")
    assert full_pdf.headers["cache-control"] == "no-store"

    full_excel = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/full-report.xlsx",
        headers=admin_headers,
    )
    assert full_excel.status_code == 200, full_excel.text
    workbook = load_workbook(BytesIO(full_excel.content), read_only=True)
    assert "Cari Özet" in workbook.sheetnames
    assert "Finans İşlemleri" in workbook.sheetnames
    assert "OSGB Sözleşmeleri" in workbook.sheetnames
    assert "Belge Envanteri" in workbook.sheetnames

    workplace_headers = {"Authorization": f"Bearer {_token(client, seed['users'][2], seed['password'])}"}
    workplace_excel = client.get(
        f"/api/v1/companies/{seed['company_1']}/status/full-report.xlsx",
        headers=workplace_headers,
    )
    assert workplace_excel.status_code == 200, workplace_excel.text
    workplace_workbook = load_workbook(BytesIO(workplace_excel.content), read_only=True)
    assert "Cari Özet" not in workplace_workbook.sheetnames
    assert "Finans İşlemleri" not in workplace_workbook.sheetnames
    assert "OSGB Sözleşmeleri" not in workplace_workbook.sheetnames


def test_forgot_password_keeps_neutral_response_and_does_not_crash(client):
    seed = _seed()
    response = client.post("/api/v1/auth/forgot-password", json={"email": seed["users"][1]})
    assert response.status_code == 200, response.text
    assert response.json() == {
        "message": "Eğer hesap varsa sıfırlama bağlantısı e-posta ile gönderildi."
    }

    unknown = client.post("/api/v1/auth/forgot-password", json={"email": "unknown@example.com"})
    assert unknown.status_code == 200
    assert unknown.json() == response.json()


def test_report_exports_neutralize_excel_formula_and_escape_pdf_markup():
    from app.services.workplace_status_reports import (
        build_workplace_status_excel,
        build_workplace_status_pdf,
    )

    payload = {
        "company": {"name": "<b>Deneme & İşyeri</b>", "sgk_registry_no": "=1+1"},
        "status_center": {
            "overall_label": "İzleme",
            "completion_pct": 50,
            "generated_at": "2026-08-04T00:00:00Z",
            "items": [{
                "code": "risk",
                "title": "=HYPERLINK(\"https://evil.invalid\")",
                "status_label": "Eksik",
                "detail": "<script>değil</script> & güvenli",
                "responsible_role": "+Sorumlu",
                "source": "risk_assessments",
                "module": "risk",
                "critical": True,
            }],
            "deadlines": [],
        },
    }
    excel = build_workplace_status_excel(payload)
    workbook = load_workbook(BytesIO(excel), data_only=False)
    assert workbook["Durum Özeti"]["B2"].value == "'=1+1"
    assert workbook["Süreç Durumları"]["B2"].value.startswith("'=")
    assert workbook["Süreç Durumları"]["E2"].value.startswith("'+")

    pdf = build_workplace_status_pdf(payload)
    assert pdf.startswith(b"%PDF")


def test_capa_board_status_and_excel_share_full_company_scope(client):
    seed = _seed()
    from app.core.database import SessionLocal
    from app.models.entities import Company, Hazard, HazardCategory, RiskAssessment, RiskDof, IncidentEvent, IncidentDof

    today = date.today()
    with SessionLocal() as db:
        own_company = db.get(Company, seed["company_1"])
        sibling = Company(name="Aynı OSGB farklı firma", osgb_id=own_company.osgb_id, is_active=True)
        category = HazardCategory(name="DÖF kapsam testi", sort_order=1)
        db.add_all([sibling, category])
        db.flush()
        sibling_id = sibling.id
        hazard = Hazard(category_id=category.id, code="H-CAPA", name="Test tehlikesi")
        db.add(hazard)
        db.flush()
        risks = [RiskAssessment(
            risk_code=f"R-CAPA-{index}", company_id=seed["company_1"], hazard_id=hazard.id,
            activity="Test faaliyeti", risk_definition="Test riski", probability=2, severity=2,
            risk_score=4, risk_level="Düşük", created_by_id=seed["admin_id"],
        ) for index in range(1016)]
        foreign_risk = RiskAssessment(
            risk_code="R-SIBLING", company_id=sibling_id, hazard_id=hazard.id,
            activity="Diğer firma", risk_definition="Kapsam dışı", probability=2, severity=2,
            risk_score=4, risk_level="Düşük", created_by_id=seed["admin_id"],
        )
        db.add_all([*risks, foreign_risk])
        db.flush()
        for index, risk in enumerate([*risks, foreign_risk]):
            db.add(RiskDof(
                dof_code=f"D-CAPA-{index}", risk_id=risk.id, description="Test önlemi",
                term_date=today-timedelta(days=1), is_completed=index==1015,
                created_by_id=seed["admin_id"],
            ))
        for index, (cid, status, term) in enumerate([
            (seed["company_1"], "Açık", today-timedelta(days=1)),
            (seed["company_1"], "closed", today-timedelta(days=1)),
            (seed["company_1"], "Açık", None),
            (sibling_id, "Açık", today-timedelta(days=1)),
            (seed["company_2"], "Açık", today-timedelta(days=1)),
        ]):
            incident = IncidentEvent(
                form_no=f"I-CAPA-{index}", company_id=cid, event_type="ramak_kala",
                event_date=today, short_summary="Test olayı", created_by_id=seed["admin_id"],
            )
            db.add(incident)
            db.flush()
            db.add(IncidentDof(
                dof_no=f"ID-CAPA-{index}", incident_id=incident.id, finding="Test bulgusu",
                status=status, term_date=term, created_by_id=seed["admin_id"],
            ))
        db.commit()

    headers = {"Authorization": f"Bearer {_token(client, seed['users'][0], seed['password'])}"}
    response = client.get(f"/api/v1/incidents/capa-board?company_id={seed['company_1']}", headers=headers)
    assert response.status_code == 200, response.text
    board = response.json()
    assert board["summary"] == {"total": 1019, "open": 1017, "completed": 2, "overdue": 1016}
    assert len(board["items"]) == 1019
    assert {row["company_id"] for row in board["items"]} == {seed["company_1"]}
    assert "D-CAPA-1015" in {row["code"] for row in board["items"]}
    assert not {"D-CAPA-1016", "ID-CAPA-3", "ID-CAPA-4"}.intersection(row["code"] for row in board["items"])

    status = client.get(f"/api/v1/companies/{seed['company_1']}/status", headers=headers).json()
    items = {row["code"]: row for row in status["status_center"]["items"]}
    assert items["capa"]["count"] == board["summary"]["open"]
    assert items["capa"]["detail"] == "1019 DÖF kaydı; 1017 açık DÖF; 1016 gecikmiş."
    assert status["counts"]["open_dofs"] == 1017
    assert status["counts"]["overdue_dofs"] == 1016
    assert items["risk_assessment_validity"]["status"] == "attention"
    assert items["risk_assessment_validity"]["status_label"] == "Tarih eksik"
    assert not any(row["source"] == "Risk Değerlendirmesi" for row in status["status_center"]["deadlines"])

    exported = client.get(f"/api/v1/incidents/capa-board.xlsx?company_id={seed['company_1']}", headers=headers)
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    exported_codes = {row[1] for row in workbook.active.iter_rows(values_only=True) if row[0] in ("Risk", "Olay")}
    assert exported_codes == {row["code"] for row in board["items"]}
    assert client.get("/api/v1/incidents/capa-board", headers=headers).status_code == 422
    assert client.get(f"/api/v1/incidents/capa-board?company_id={seed['company_2']}", headers=headers).status_code == 403
    workplace_headers = {"Authorization": f"Bearer {_token(client, seed['users'][2], seed['password'])}"}
    assert client.get(f"/api/v1/incidents/capa-board?company_id={sibling_id}", headers=workplace_headers).status_code == 403
    assert client.get(f"/api/v1/incidents/capa-board.xlsx?company_id={sibling_id}", headers=workplace_headers).status_code == 403


def _seed_dof_actions(seed):
    from app.core.database import SessionLocal
    from app.core.security import get_password_hash
    from app.models.entities import (HazardCategory, Hazard, RiskAssessment, RiskDof, IncidentEvent,
                                    IncidentDof, User, UserRole, Company, IsgProfessional,
                                    ProfessionalType, WorkplaceAssignment, AssignmentStatus)
    with SessionLocal() as db:
        osgb_id = db.get(Company, seed['company_1']).osgb_id
        doctor = User(email='dof-doctor@example.com', full_name='Test Hekimi',
                      hashed_password=get_password_hash(seed['password']), role=UserRole.WORKPLACE_PHYSICIAN,
                      osgb_id=osgb_id, is_active=True)
        professional = IsgProfessional(osgb_id=osgb_id, full_name='Test Hekimi', email=doctor.email,
                                      professional_type=ProfessionalType.WORKPLACE_PHYSICIAN,
                                      certificate_number='DOF-DOC-1', is_active=True)
        category = HazardCategory(name='DÖF işlem testi', sort_order=1)
        db.add_all([doctor, professional, category]); db.flush()
        db.add(WorkplaceAssignment(osgb_id=osgb_id, company_id=seed['company_1'], professional_id=professional.id,
                                  professional_type=ProfessionalType.WORKPLACE_PHYSICIAN,
                                  start_date=date.today(), status=AssignmentStatus.ACTIVE))
        hazard = Hazard(category_id=category.id, code='DOF-ACTION', name='Koruyucu eksikliği')
        db.add(hazard); db.flush()
        actions = []
        for index, company_id in enumerate([seed['company_1'], seed['company_2']]):
            risk = RiskAssessment(risk_code=f'R-ACTION-{index}', company_id=company_id, hazard_id=hazard.id,
                                  activity='Makine çalışması', risk_definition='Koruyucu eksikliği',
                                  probability=2, severity=2, risk_score=4, risk_level='Düşük',
                                  created_by_id=seed['admin_id'])
            incident = IncidentEvent(form_no=f'I-ACTION-{index}', company_id=company_id, event_type='ramak_kala',
                                     event_date=date.today(), short_summary='Koruyucu eksikliği tespit edildi.',
                                     created_by_id=seed['admin_id'])
            db.add_all([risk, incident]); db.flush()
            dof = RiskDof(dof_code=f'D-ACTION-{index}', risk_id=risk.id, description='Makine koruyucusu takılacak.',
                          term_date=date.today()-timedelta(days=3), created_by_id=seed['admin_id'])
            incident_dof = IncidentDof(dof_no=f'ID-ACTION-{index}', incident_id=incident.id,
                                      finding='Makine koruyucusu eksik.', corrective_action='Makine koruyucusu takılacak.',
                                      preventive_action='Periyodik kontrol yapılacak.', responsible_person='Test Sorumlusu',
                                      term_date=date.today()-timedelta(days=3), created_by_id=seed['admin_id'])
            db.add_all([dof, incident_dof]); db.flush()
            actions.append({'risk_id': risk.id, 'dof_id': dof.id, 'incident_id': incident.id, 'incident_dof_id': incident_dof.id})
        db.commit()
        return actions, doctor.email


@pytest.mark.parametrize('role', ['workplace', 'specialist', 'physician'])
def test_dof_actions_allow_only_authorized_workplaces_and_update_counts(client, role):
    seed = _seed()
    (own, foreign), doctor = _seed_dof_actions(seed)
    email = {'workplace': seed['users'][2], 'specialist': seed['users'][3], 'physician': doctor}[role]
    headers = {'Authorization': f"Bearer {_token(client, email, seed['password'])}"}
    base = f"/api/v1/risks/{own['risk_id']}/dofs/{own['dof_id']}"
    incident_base = f"/api/v1/incidents/{own['incident_id']}/dofs/{own['incident_dof_id']}"
    for path in [f"/api/v1/risks/{foreign['risk_id']}/dofs/{foreign['dof_id']}",
                 f"/api/v1/incidents/{foreign['incident_id']}/dofs/{foreign['incident_dof_id']}"]:
        assert client.patch(path, headers=headers, json={'responsible_person': 'Test Sorumlusu'}).status_code == 403
        assert client.post(path+'/complete', headers=headers, json={}).status_code == 403
    response = client.patch(base, headers=headers, json={'responsible_person': 'Bakım Müdürü', 'term_date': date.today().isoformat()})
    assert response.status_code == 200, response.text
    response = client.patch(incident_base, headers=headers, json={'corrective_action': 'Makine koruyucusu yenilenecek.'})
    assert response.status_code == 200, response.text
    assert client.patch(f"/api/v1/risks/{own['risk_id']}/dofs/{foreign['dof_id']}", headers=headers, json={'responsible_person': 'Test Sorumlusu'}).status_code == 404
    completed_on = (date.today()-timedelta(days=1)).isoformat()
    response = client.post(base+'/complete', headers=headers, json={'completion_date': completed_on, 'completion_note': 'Makine koruyucusu takıldı ve kontrol edildi.'})
    assert response.status_code == 200, response.text
    assert response.json()['is_completed'] is True
    assert response.json()['completion_date'] == completed_on
    response = client.post(incident_base+'/complete', headers=headers, json={'completion_date': completed_on, 'effectiveness_note': 'Makine koruyucusu takıldı ve kontrol edildi.', 'close_approval': 'Test Sorumlusu'})
    assert response.status_code == 200, response.text
    board = client.get(f"/api/v1/incidents/capa-board?company_id={seed['company_1']}", headers=headers).json()
    assert board['summary'] == {'total': 2, 'open': 0, 'completed': 2, 'overdue': 0}
    assert all(row['parent_id'] and row['completion_date'] == completed_on for row in board['items'])
    # Repeated clicks/retries must not overwrite the recorded completion date.
    assert client.post(base+'/complete', headers=headers, json={}).json()['completion_date'] == completed_on
    if role != 'specialist':
        assert client.patch(f"/api/v1/risks/{own['risk_id']}", headers=headers, json={'probability': 1}).status_code == 403


def test_osgb_manager_and_read_only_user_cannot_create_edit_or_complete_dofs(client):
    seed = _seed()
    (own, _), _doctor = _seed_dof_actions(seed)
    for email in seed['users'][:2]:
        headers = {'Authorization': f"Bearer {_token(client, email, seed['password'])}"}
        assert client.get(f"/api/v1/incidents/capa-board?company_id={seed['company_1']}", headers=headers).status_code == 200
        for source, parent, identifier in [('risks', own['risk_id'], own['dof_id']), ('incidents', own['incident_id'], own['incident_dof_id'])]:
            base = f'/api/v1/{source}/{parent}/dofs'
            assert client.post(base, headers=headers, json={}).status_code == 403
            assert client.patch(f'{base}/{identifier}', headers=headers, json={'responsible_person': 'Test Sorumlusu'}).status_code == 403
            assert client.post(f'{base}/{identifier}/complete', headers=headers, json={}).status_code == 403


def test_continuous_deadline_migration_preserves_manual_mixed_and_completed_actions(client):
    from datetime import datetime
    import importlib.util
    from pathlib import Path
    from sqlalchemy import select
    from app.core.database import SessionLocal
    from app.models.entities import RiskAssessment, RiskDof, AuditLog
    seed = _seed()
    (own, _), _doctor = _seed_dof_actions(seed)
    created = datetime(2026, 9, 23, 10, 0)
    base_date = datetime(2026, 6, 9)
    due = base_date.date() + timedelta(days=90)
    ids = {}
    with SessionLocal() as db:
        hazard_id = db.get(RiskAssessment, own['risk_id']).hazard_id
        for index, mode in enumerate(['continuous', 'mixed', 'edited', 'completed', 'manual']):
            fp = f'fingerprint-{index}'
            risk = RiskAssessment(risk_code=f'R-CONT-{index}', company_id=seed['company_1'], hazard_id=hazard_id,
                                  activity='Kontrol', risk_definition='Kayma riski', probability=2, severity=2,
                                  risk_score=4, risk_level='Düşük', created_by_id=seed['admin_id'], record_origin='excel_import',
                                  source_fingerprint=fp, observed_at=base_date, created_at=created, updated_at=created,
                                  term_text='30 gün / Sürekli izleme' if mode == 'mixed' else 'Sürekli izleme',
                                  term_date=due, term_days=90, term_suggested=90, term_overridden=False)
            db.add(risk); db.flush()
            dof = RiskDof(dof_code=f'D-CONT-{index}', risk_id=risk.id, description='Düzenli kontrol yap.',
                          client_reference=None if mode == 'manual' else f'excel:{fp}:dof', term_date=due,
                          created_by_id=seed['admin_id'], created_at=created,
                          updated_at=created+timedelta(days=1) if mode == 'edited' else created,
                          is_completed=mode == 'completed', status='Tamamlandı' if mode == 'completed' else 'Açık')
            db.add(dof); db.flush(); ids[mode] = dof.id
        db.commit()
        spec = importlib.util.spec_from_file_location('deadline_migration', Path(__file__).parents[1]/'alembic/versions/0129_continuous_risk_deadlines.py')
        migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(migration)
        changed = migration.repair_continuous_deadlines(db.connection())
        assert changed['dofs'] == 1
        assert migration.repair_continuous_deadlines(db.connection()) == {'risks': 0, 'dofs': 0}
        db.commit(); db.expire_all()
        assert db.get(RiskDof, ids['continuous']).term_date is None
        for mode in ['mixed', 'edited', 'completed', 'manual']:
            assert db.get(RiskDof, ids[mode]).term_date == due
        assert db.scalar(select(AuditLog).where(AuditLog.entity_type == 'risk_dof', AuditLog.entity_id == str(ids['continuous']))) is not None
    headers = {'Authorization': f"Bearer {_token(client, seed['users'][2], seed['password'])}"}
    board = client.get(f"/api/v1/incidents/capa-board?company_id={seed['company_1']}", headers=headers).json()
    continuous = next(row for row in board['items'] if row['key'] == f"r-{ids['continuous']}")
    assert continuous['term_kind'] == 'continuous'
    assert continuous['is_overdue'] is False
