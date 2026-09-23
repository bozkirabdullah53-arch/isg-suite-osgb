from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.services.risk_excel_import import build_import_risk_code, parse_risk_workbook


def test_import_risk_code_does_not_depend_on_visible_global_risk_count() -> None:
    first = build_import_risk_code(174, "a" * 32)
    same_row = build_import_risk_code(174, "a" * 32)
    other_company = build_import_risk_code(124, "a" * 32)
    collision = build_import_risk_code(174, "a" * 32, 1)

    assert first == same_row
    assert first != other_company
    assert first != collision
    assert first.startswith("RSK-")
    assert len(first) == 20
    assert len(collision) == 20


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    main = workbook.active
    main.title = "Yazdırmaya Hazır Rapor"
    main.append(["ERDİL AKÜ RİSK ANALİZİ", None, None])
    main.append(["Rapor Tarihi: 09.06.2026", None, None])
    main.append(
        [
            "PN",
            "Proses / Alan",
            "Faaliyet",
            "Risk Unsuru",
            "Tehlike",
            "Olası Sonuç",
            "O (1-5)",
            "Ş (1-5)",
            "RP",
            "Mevcut Önlemler",
            "Alınması Gereken İlave Önlemler",
            "Termin süresi",
            "Sorumlu",
            "İlgili Mevzuat Dayanağı",
        ]
    )
    main.append(
        [
            1,
            "Kantar",
            "Araç tartım",
            "Araç-yaya etkileşimi",
            "Ağır araç hareketi",
            "Ezilme",
            3,
            5,
            15,
            "Uyarı levhası",
            "Yaya yolu ayrılmalı",
            "7 gün",
            "İSG Uzmanı",
            "6331 m.5",
        ]
    )
    # Page-signature text below the table must not be treated as a bad risk.
    main.append([None, None, None, "İşyeri Hekimi", None, None, None, None])
    main.append(["PN", "Proses / Alan", "Faaliyet", "Risk Unsuru", "Tehlike", "Olası Sonuç", "O (1-5)", "Ş (1-5)", "RP"])
    main.append(
        [
            2,
            "Depo",
            "İstifleme",
            "Devrilme",
            "Dengesiz istif",
            "Yaralanma",
            2,
            4,
            99,
            "Raf kontrolü",
            "İstif talimatı uygulanmalı",
            "30 gün",
            "Depo Sorumlusu",
            "6331 m.10",
        ]
    )

    appendix = workbook.create_sheet("Ek Fotoğraflı Riskler")
    appendix.append(["Ek fotoğraflı riskler"])
    appendix.append(
        [
            "No",
            "Foto No",
            "Proses / Alan",
            "Faaliyet",
            "Risk Unsuru",
            "Tehlike",
            "Olası Sonuç",
            "O",
            "Ş",
            "RP",
            "Mevcut Önlemler",
            "Alınması Gereken İlave Önlemler",
            "Termin süresi",
            "Sorumlu",
            "İlgili Mevzuat Dayanağı",
        ]
    )
    appendix.append(
        [
            3,
            20,
            "Makine",
            "Bakım",
            "Hareketli parça",
            "Sıkışma",
            "El yaralanması",
            3,
            4,
            12,
            "Muhafaza",
            "LOTO uygulanmalı",
            "7 gün",
            "Bakım Sorumlusu",
            "6331 m.10",
        ]
    )

    tracking = workbook.create_sheet("Düzeltici Faaliyet Takip")
    tracking.append(["Aksiyon ID", "Proses / Alan", "Faaliyet", "Risk Unsuru", "Tehlike", "O", "Ş"])
    tracking.append(["DÖF-1", "Kantar", "Araç tartım", "Araç", "Çarpma", 3, 5])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parser_keeps_canonical_and_photo_risk_rows_only() -> None:
    result = parse_risk_workbook(_workbook_bytes(), filename="risk.xlsx")

    assert len(result["rows"]) == 3
    assert {row["source_sheet"] for row in result["rows"]} == {
        "Yazdırmaya Hazır Rapor",
        "Ek Fotoğraflı Riskler",
    }
    assert result["metadata"]["assessment_date"] == "2026-06-09"
    assert result["rows"][0]["probability"] == 3
    assert result["rows"][0]["severity"] == 5
    assert result["rows"][0]["source_score"] == 15
    assert result["warnings"]
    assert not result["errors"]


def test_parser_fingerprint_includes_source_number_and_preserves_fields() -> None:
    result = parse_risk_workbook(_workbook_bytes())
    first, second, appendix = result["rows"]

    assert first["source_pn"] == "1"
    assert second["source_pn"] == "2"
    assert appendix["photo_no"] == "20"
    assert first["legislation_basis"] == "6331 m.5"
    assert first["term_days_hint"] == 7
    assert len({row["fingerprint"] for row in result["rows"]}) == 3
